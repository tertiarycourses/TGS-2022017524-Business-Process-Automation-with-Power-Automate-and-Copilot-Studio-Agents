#!/usr/bin/env python3
"""LG v7.3 — mirror the new LEARNER-GUIDE.md content into the LG DOCX.

Reads the Lab 4b and Lab 5b sections and the Lab 4 skill-upload steps straight
out of LEARNER-GUIDE.md and renders them into the DOCX at the matching anchors,
so the Markdown mirror and the DOCX cannot diverge. Also bumps the version
header and adds the 7.3 Document Version Control Record row.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "LEARNER-GUIDE.md"
DOCX = ROOT / "courseware" / "LG-Business Process Automation with Power Automate and Copilot Studio Agents.docx"

md_text = MD.read_text(encoding="utf-8")


def extract(start_pat, end_pat):
    m = re.search(start_pat, md_text)
    assert m, start_pat
    rest = md_text[m.start():]
    e = re.search(end_pat, rest[1:])
    assert e, end_pat
    return rest[:e.start() + 1]


SEC_4B = extract(r"### Lab 4b — Multi-Agent Content Team \(optional\)",
                 r"\n## Day 2 — Agent Flows, Human Review and RAG")
SEC_5B = extract(r"### Lab 5b — Calling Workflow from Agent",
                 r"\n### Module 4: Agent Flows, HTTP and the Boundary of Agency")
SEC_SKILL = extract(r"\*\*Uploading a skill package — step by step\.\*\*",
                    r"\nThree things follow")

doc = Document(DOCX)


# --------------------------------------------------------------- renderer

def add_runs(p, textline):
    """Render **bold**, `code`, *italic*, [text](url) inline markdown into runs."""
    textline = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", textline)
    pos = 0
    pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")
    for m in pattern.finditer(textline):
        if m.start() > pos:
            p.add_run(textline[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**"):
            r = p.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("`"):
            r = p.add_run(tok[1:-1])
            r.font.name = "Consolas"
        else:
            r = p.add_run(tok[1:-1])
            r.italic = True
        pos = m.end()
    if pos < len(textline):
        p.add_run(textline[pos:])


def render_blocks(md_section):
    """Yield ready DOCX block elements appended to the document end."""
    made = []
    lines = md_section.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line or line == "---":
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_runs(p, line[4:])
            made.append(p)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"-{3,}", c) for c in cells):
                    rows.append(cells)
                i += 1
            i -= 1
            ncols = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncols)
            t.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci in range(ncols):
                    cell = t.cell(ri, ci)
                    cell.paragraphs[0].text = ""
                    add_runs(cell.paragraphs[0], row[ci] if ci < len(row) else "")
                    if ri == 0:
                        for r in cell.paragraphs[0].runs:
                            r.bold = True
                    for r in cell.paragraphs[0].runs:
                        r.font.size = Pt(9.5)
            made.append(t)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:])
            made.append(p)
        elif re.match(r"^\d+\. ", line):
            num, body = line.split(". ", 1)
            p = doc.add_paragraph()
            p.add_run(f"{num}.  ")
            add_runs(p, body.strip())
            made.append(p)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            r = p.add_run("Note: ")
            r.bold = True
            add_runs(p, line[2:])
            made.append(p)
        elif line.startswith("!["):
            m = re.match(r"!\[([^\]]*)\]\(<?([^)>]+)>?\)", line)
            alt, path = m.group(1), ROOT / m.group(2)
            if path.exists():
                doc.add_picture(str(path), width=Inches(6.0))
                made.append(doc.paragraphs[-1])
            cap = doc.add_paragraph()
            r = cap.add_run(f"Figure: {alt}")
            r.italic = True
            made.append(cap)
        elif re.fullmatch(r"\*[^*].*\*", line):
            p = doc.add_paragraph()
            r = p.add_run(line[1:-1])
            r.italic = True
            made.append(p)
        else:
            # merge hard-wrapped continuation lines of the same paragraph
            buf = [line]
            while (i + 1 < len(lines) and lines[i + 1].strip()
                   and not re.match(r"^(#{1,6} |\||- |\d+\. |> |!\[|\*\*[^*]|---)", lines[i + 1].strip())):
                i += 1
                buf.append(lines[i].strip())
            p = doc.add_paragraph()
            add_runs(p, " ".join(buf))
            made.append(p)
        i += 1
    return made


def insert_before(anchor_el, elements):
    for obj in elements:
        el = obj._p if hasattr(obj, "_p") else obj._tbl
        anchor_el.addprevious(el)


def find_para(pred, start=0):
    for idx, p in enumerate(doc.paragraphs[start:], start):
        if pred(p):
            return idx, p
    raise SystemExit("anchor paragraph not found")


# 1. version header + record row
_, vp = find_para(lambda p: p.text.strip() == "Version 7.2")
vp.runs[0].text = "Version 7.3"
vt = doc.tables[0]
row = vt.add_row()
for ci, val in enumerate([
    "7.3", "6 Aug 2026",
    "Aligned to the 13-lab structure and the 80-slide v7.3 deck — added Lab 4b "
    "(Multi-Agent Content Team) and Lab 5b (Calling Workflow from Agent) as full "
    "activities, and expanded Lab 4 with detailed step-by-step skill-package "
    "upload instructions.",
    "Course Development Team",
]):
    row.cells[ci].paragraphs[0].text = val
    for r in row.cells[ci].paragraphs[0].runs:
        r.font.size = Pt(9)
print("version bumped to 7.3")

# 2. skill-upload steps replace the one-line "To add one:" paragraph
_, skill_anchor = find_para(lambda p: p.text.strip().startswith("To add one:"))
insert_before(skill_anchor._p, render_blocks(SEC_SKILL))
skill_anchor._p.getparent().remove(skill_anchor._p)
print("skill upload steps inserted")

# 3. Lab 4's Next pointer + the Lab 4b section before the Day 2 heading
_, next4 = find_para(lambda p: p.text.strip() == "Next: Lab 5 — Calling Agent from Workflow"
                     or p.text.strip() == "Next: Lab 5 — Calling Agent from Workflow")
for r in next4.runs:
    r.text = ""
r0 = next4.runs[0] if next4.runs else next4.add_run()
r0.text = "Next: Lab 4b — Multi-Agent Content Team (optional), then Lab 5 — Calling Agent from Workflow"
r0.bold = True

_, day2 = find_para(lambda p: p.style.name == "Heading 1" and p.text.strip().startswith("Day 2 —"))
insert_before(day2._p, render_blocks(SEC_4B))
print("Lab 4b section inserted")

# 4. Lab 5b section before the Module 4 heading
_, mod4 = find_para(lambda p: p.style.name == "Heading 2" and p.text.strip().startswith("Module 4:"))
insert_before(mod4._p, render_blocks(SEC_5B))
print("Lab 5b section inserted")

doc.save(DOCX)
print("saved", DOCX.name)
