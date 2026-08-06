#!/usr/bin/env python3
"""v7.3 — align the deck to the thirteen-lab structure and the new agent UI.

Four new slides:
  Agent Settings                       after "Multiple Agents — A Split Is a Governance Decision"
  Publishing an Agent — Channels       after "Publishing to Teams"
  Lab 4b — Multi-Agent Content Team    after the Lab 4 workflow slide
  Lab 5b — Calling Workflow from Agent after the Lab 5 workflow slide

Plus: the lab journey rebuilt for 13 labs, schedule rows folding in 4b/5b, the
Module 3 divider's Applied-in line, the cover version, and footer renumbering.
Both dialog slides reproduce the live Copilot Studio dialogs (user screenshots,
6 Aug 2026): the publish-channels dialog and the Agent settings AI & behavior
page.
"""

from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.packuri import PackURI
from pptx.parts.presentation import PresentationPart
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
CW = ROOT / "courseware"
TITLE = "Business Process Automation with Power Automate and Copilot Studio Agents"
SRC = CW / f"{TITLE}-v7.2.pptx"
DECK = CW / f"{TITLE}-v7.3.pptx"
FLOW_5B = ROOT / "labs/Lab 5b - Calling Workflow from Agent/assets/flowchart.png"

BLUE = RGBColor(0x1F, 0x6F, 0xEB)
TEAL = RGBColor(0x10, 0x8A, 0x73)
GREEN = RGBColor(0x16, 0x84, 0x5B)
AMBER = RGBColor(0xC7, 0x76, 0x00)
VIOLET = RGBColor(0x6D, 0x3F, 0xD2)
RED = RGBColor(0xC2, 0x41, 0x3A)
INK = RGBColor(0x16, 0x1B, 0x26)
GREY = RGBColor(0x5B, 0x63, 0x72)
LINE = RGBColor(0xD7, 0xE0, 0xEA)
LIGHT = RGBColor(0xF5, 0xF8, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FOOTER_TEXT = (
    "Business Process Automation with Power Automate & Copilot Studio Agents  ·  "
    "TGS-2022017524  ·  © 2026 Tertiary Infotech Academy Pte Ltd"
)

prs = Presentation(SRC)
SW = prs.slide_width


def box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True, line_w=1.2):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_w)
    shape.shadow.inherit = False
    return shape


def rect(slide, x, y, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def text(slide, x, y, w, h, value, size=18, color=INK, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=0.08, space=0):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    for i, line_text in enumerate(str(value).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.space_after = Pt(space)
        r = p.add_run()
        r.text = line_text
        r.font.name = "Arial"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return tb


def title(slide, heading, kicker):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    rect(slide, 0, 0, Inches(0.22), Inches(1.48), BLUE)
    text(slide, Inches(0.72), Inches(0.34), Inches(11.8), Inches(0.28), kicker.upper(), 12, BLUE, True)
    text(slide, Inches(0.72), Inches(0.68), Inches(11.8), Inches(0.66), heading, 29, INK, True)
    rect(slide, Inches(0.72), Inches(1.5), Inches(11.85), Inches(0.015), LINE)


def footer(slide, slide_no):
    text(slide, Inches(0.72), Inches(6.95), Inches(11.85), Inches(0.24), FOOTER_TEXT, 8, GREY)
    text(slide, Inches(0.72), Inches(6.95), Inches(11.85), Inches(0.24), str(slide_no), 8, GREY, False, PP_ALIGN.RIGHT)


def takeaway(slide, y, message, color=BLUE):
    box(slide, Inches(0.72), Inches(y), Inches(11.85), Inches(0.62), LIGHT, color, line_w=1.6)
    text(slide, Inches(1.0), Inches(y), Inches(11.3), Inches(0.62), message, 14, INK, True)


def mini_card(slide, x, y, w, h, label, body, accent, label_size=12, body_size=9.5):
    box(slide, x, y, w, h, LIGHT, LINE)
    rect(slide, x, y, Inches(0.09), h, accent)
    text(slide, x + Inches(0.22), y + Inches(0.05), w - Inches(0.4), Inches(0.3), label, label_size, accent, True)
    text(slide, x + Inches(0.22), y + Inches(0.36), w - Inches(0.4), h - Inches(0.44), body, body_size, GREY,
         anchor=MSO_ANCHOR.TOP, space=2)


def lab_banner(slide, num, heading, body):
    y = 6.08
    box(slide, Inches(0.72), Inches(y), Inches(11.85), Inches(0.78), WHITE, GREEN, line_w=1.8)
    tag = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.92), Inches(y + 0.14), Inches(1.15), Inches(0.5))
    tag.fill.solid()
    tag.fill.fore_color.rgb = GREEN
    tag.line.fill.background()
    tag.shadow.inherit = False
    text(slide, Inches(0.92), Inches(y + 0.14), Inches(1.15), Inches(0.5), f"LAB {num}", 13, WHITE, True, PP_ALIGN.CENTER)
    text(slide, Inches(2.25), Inches(y + 0.02), Inches(10.1), Inches(0.36), heading, 14, INK, True)
    text(slide, Inches(2.25), Inches(y + 0.38), Inches(10.1), Inches(0.34), body, 11, GREY)


def toggle(slide, x, y, on=True):
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(0.44), Inches(0.22))
    pill.adjustments[0] = 0.5
    pill.fill.solid()
    pill.fill.fore_color.rgb = BLUE if on else GREY
    pill.line.fill.background()
    pill.shadow.inherit = False
    knob = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, x + (Inches(0.25) if on else Inches(0.03)), y + Inches(0.03), Inches(0.16), Inches(0.16))
    knob.fill.solid()
    knob.fill.fore_color.rgb = WHITE
    knob.line.fill.background()
    knob.shadow.inherit = False


def checkbox(slide, x, y):
    bx = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(0.2), Inches(0.2))
    bx.fill.solid()
    bx.fill.fore_color.rgb = BLUE
    bx.line.fill.background()
    bx.shadow.inherit = False
    text(slide, x - Inches(0.02), y - Inches(0.03), Inches(0.24), Inches(0.24), "✓", 10, WHITE, True,
         PP_ALIGN.CENTER, margin=0)


# --------------------------------------------------------------- new slides

def draw_agent_settings(s):
    text(s, Inches(0.72), Inches(1.60), Inches(11.85), Inches(0.32),
         "Every agent has a Settings page. Two of its switches decide how this course's agents behave — "
         "and neither takes effect until you save and publish.", 13, GREY, anchor=MSO_ANCHOR.TOP)

    # ---- the dialog, as it appears in the product ----
    dx, dy, dw, dh = Inches(0.72), Inches(2.10), Inches(7.30), Inches(3.72)
    box(s, dx, dy, dw, dh, WHITE, LINE, line_w=1.4)
    text(s, dx + Inches(0.14), dy + Inches(0.06), Inches(3.0), Inches(0.3), "Agent settings", 13, INK, True)
    box(s, dx + Inches(0.14), dy + Inches(0.42), dw - Inches(0.28), Inches(0.32), LIGHT, LIGHT)
    text(s, dx + Inches(0.26), dy + Inches(0.42), dw - Inches(0.5), Inches(0.32),
         "ⓘ  Changes to settings take effect after you save and publish the agent.", 9, GREY)

    nav_items = ["Agent details", "AI & behavior", "Safety & access", "Greeting & prompts"]
    ny = dy + Inches(0.92)
    for i, item in enumerate(nav_items):
        iy = ny + Inches(0.46) * i
        if item == "AI & behavior":
            box(s, dx + Inches(0.14), iy, Inches(2.05), Inches(0.38), LIGHT, LINE, line_w=0.75)
            rect(s, dx + Inches(0.14), iy, Inches(0.05), Inches(0.38), BLUE)
            text(s, dx + Inches(0.30), iy, Inches(1.85), Inches(0.38), item, 10, INK, True)
        else:
            text(s, dx + Inches(0.30), iy, Inches(1.85), Inches(0.38), item, 10, GREY)

    px = dx + Inches(2.45)
    pw = dw - Inches(2.65)
    text(s, px, dy + Inches(0.88), pw, Inches(0.28), "AI & behavior", 12, INK, True)
    text(s, px, dy + Inches(1.14), pw, Inches(0.26),
         "Tune how the agent responds, reasons, and communicates.", 8.5, GREY)

    text(s, px, dy + Inches(1.46), pw, Inches(0.26), "Orchestration", 10.5, INK, True)
    oy = dy + Inches(1.74)
    box(s, px, oy, pw, Inches(0.78), WHITE, LINE, line_w=1.0)
    text(s, px + Inches(0.12), oy + Inches(0.05), pw - Inches(0.8), Inches(0.28),
         "Allow other agents to connect", 10, INK, True)
    text(s, px + Inches(0.12), oy + Inches(0.33), pw - Inches(0.8), Inches(0.42),
         "Let other agents in your organization invoke this agent as a tool.", 8.5, GREY, anchor=MSO_ANCHOR.TOP)
    toggle(s, px + pw - Inches(0.6), oy + Inches(0.10), on=True)

    text(s, px, dy + Inches(2.62), pw, Inches(0.26), "Safety", 10.5, INK, True)
    sy = dy + Inches(2.90)
    box(s, px, sy, pw, Inches(0.68), WHITE, LINE, line_w=1.0)
    text(s, px + Inches(0.12), sy + Inches(0.04), pw - Inches(1.3), Inches(0.28), "Moderation level", 10, INK, True)
    text(s, px + Inches(0.12), sy + Inches(0.30), pw - Inches(1.3), Inches(0.34),
         "Controls how strictly responses are filtered for unsafe content.", 8.5, GREY, anchor=MSO_ANCHOR.TOP)
    box(s, px + pw - Inches(1.12), sy + Inches(0.17), Inches(0.98), Inches(0.34), WHITE, LINE, line_w=1.0)
    text(s, px + pw - Inches(1.04), sy + Inches(0.17), Inches(0.9), Inches(0.34), "Medium  ⌄", 9, INK)

    # ---- what each control means ----
    cx = Inches(8.24)
    cw = Inches(4.33)
    mini_card(s, cx, Inches(2.10), cw, Inches(1.16), "The toggle behind Labs 4 & 4b",
              "A child agent can only be invoked as a connected agent while Allow other agents to "
              "connect is on. Manager cannot delegate to a child that has it off.", VIOLET, 11)
    mini_card(s, cx, Inches(3.38), cw, Inches(1.16), "Moderation is a platform control",
              "The moderation filter runs outside the model at the level you set — a structural "
              "control, not an instruction the agent could be talked out of.", TEAL, 11)
    mini_card(s, cx, Inches(4.66), cw, Inches(1.16), "Nothing bites until publish",
              "Settings changes apply only after save and publish. A changed setting that "
              "“does nothing” is usually just unpublished.", AMBER, 11)

    lab_banner(s, "4B", "Lab 4b — Multi-Agent Content Team",
               "Marketing Manager, Research, Blog and Review  ·  Copilot Studio connected agents  ·  60–75 min · optional")


def draw_publishing_channels(s):
    text(s, Inches(0.72), Inches(1.60), Inches(11.85), Inches(0.32),
         "Publish makes a version live — a channel decides who can reach it. The publish dialog lists the "
         "channels; none of them serves your edits until you publish again.", 13, GREY, anchor=MSO_ANCHOR.TOP)

    # ---- the publish dialog ----
    dx, dy, dw, dh = Inches(0.72), Inches(2.10), Inches(7.30), Inches(3.72)
    box(s, dx, dy, dw, dh, WHITE, LINE, line_w=1.4)
    text(s, dx + Inches(0.14), dy + Inches(0.06), Inches(3.4), Inches(0.3), "Agent published", 13, INK, True)
    text(s, dx + Inches(0.14), dy + Inches(0.36), Inches(4.2), Inches(0.24),
         "✓  Last published today, 7:46 PM", 8.5, GREEN)

    channels = [
        ("Demo Website", "Shareable web link for anyone to use", False),
        ("Web app", "Embed your agent into a website", False),
        ("Teams + Microsoft 365", "Available in Copilot and Teams chats", True),
    ]
    ly = dy + Inches(0.74)
    for i, (name, sub, selected) in enumerate(channels):
        ry = ly + Inches(0.94) * i
        box(s, dx + Inches(0.14), ry, Inches(3.30), Inches(0.84),
            WHITE, BLUE if selected else LINE, line_w=1.8 if selected else 1.0)
        text(s, dx + Inches(0.28), ry + Inches(0.07), Inches(2.6), Inches(0.3), name, 10, INK, True)
        text(s, dx + Inches(0.28), ry + Inches(0.36), Inches(2.75), Inches(0.44), sub, 8.5, GREY,
             anchor=MSO_ANCHOR.TOP)
        if selected:
            checkbox(s, dx + Inches(3.12), ry + Inches(0.10))

    px = dx + Inches(3.66)
    pw = dw - Inches(3.86)
    text(s, px, dy + Inches(0.78), pw, Inches(0.3), "Teams + Microsoft 365", 12, INK, True)
    pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, dy + Inches(1.12), Inches(1.35), Inches(0.28))
    pill.adjustments[0] = 0.5
    pill.fill.solid()
    pill.fill.fore_color.rgb = GREEN
    pill.line.fill.background()
    pill.shadow.inherit = False
    text(s, px, dy + Inches(1.12), Inches(1.35), Inches(0.28), "Channel enabled", 8.5, WHITE, True, PP_ALIGN.CENTER)
    text(s, px, dy + Inches(1.50), pw, Inches(0.5),
         "Available in Copilot and Teams chats — Outlook, Word, Excel, PowerPoint and OneDrive.",
         8.5, GREY, anchor=MSO_ANCHOR.TOP)
    checkbox(s, px, dy + Inches(2.14))
    text(s, px + Inches(0.30), dy + Inches(2.08), pw - Inches(0.3), Inches(0.32),
         "Make agent available in Microsoft 365 Copilot", 9.5, INK, True)
    text(s, px, dy + Inches(2.48), pw, Inches(0.62),
         "Publishing to the Agent Store inside Microsoft 365 Copilot goes through your "
         "admin's approval.", 8.5, GREY, anchor=MSO_ANCHOR.TOP)
    text(s, px, dy + Inches(3.14), pw, Inches(0.28), "See agent in Microsoft 365  ·  See agent in Teams", 8.5, BLUE)

    # ---- teaching cards ----
    cx = Inches(8.24)
    cw = Inches(4.33)
    mini_card(s, cx, Inches(2.10), cw, Inches(1.16), "Demo Website first",
              "The shareable link is the fastest way to put the agent in a tester's hands — no "
              "Teams, no admin, anyone with the link.", BLUE, 11)
    mini_card(s, cx, Inches(3.38), cw, Inches(1.16), "Teams + Microsoft 365",
              "Enable this channel and the agent appears in Teams chat and Microsoft 365 Copilot — "
              "where colleagues already work.", VIOLET, 11)
    mini_card(s, cx, Inches(4.66), cw, Inches(1.16), "Publish, then re-publish",
              "Channels serve the published version only. Every later edit is invisible to users "
              "until you publish again.", AMBER, 11)

    lab_banner(s, "4", "Lab 4 — Agents",
               "The _deployment step publishes each agent to Microsoft Teams  ·  Copilot Studio agents  ·  45 min")


def draw_lab4b(s):
    text(s, Inches(0.72), Inches(1.62), Inches(11.85), Inches(0.34),
         "One topic, four agents: the Marketing Manager owns the conversation and delegates by stage of work — "
         "and nothing is final until a person approves.", 13, GREY, anchor=MSO_ANCHOR.TOP)

    chain = [
        ("HUMAN", "gives the topic", INK),
        ("MARKETING\nMANAGER", "orchestrates,\nnever writes", VIOLET),
        ("RESEARCH\nAGENT", "brochures + web,\nlabelled apart", BLUE),
        ("BLOG\nAGENT", "writes from the\nbrief, no web", TEAL),
        ("REVIEW\nAGENT", "checklist; did not\nwrite the draft", AMBER),
        ("HUMAN", "approves the\nfinal post", GREEN),
    ]
    cw, ch, gap = 1.82, 1.30, 0.20
    x0 = 0.72
    for i, (label, sub, color) in enumerate(chain):
        x = Inches(x0 + i * (cw + gap))
        box(s, x, Inches(2.16), Inches(cw), Inches(ch), LIGHT, color, line_w=1.6)
        text(s, x, Inches(2.24), Inches(cw), Inches(0.52), label, 11, color, True, PP_ALIGN.CENTER)
        text(s, x, Inches(2.78), Inches(cw), Inches(0.60), sub, 8.5, GREY, False, PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.TOP)
        if i:
            ar = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x0 + i * (cw + gap) - gap + 0.02),
                                    Inches(2.66), Inches(0.16), Inches(0.30))
            ar.fill.solid()
            ar.fill.fore_color.rgb = GREY
            ar.line.fill.background()
            ar.shadow.inherit = False

    mini_card(s, Inches(0.72), Inches(3.78), Inches(3.83), Inches(1.90), "Split by stage, not audience",
              "Lab 4 split agents by who may know what. This pipeline splits by what can go wrong at "
              "each stage: research, drafting and review each fail differently.", VIOLET, 12, 10)
    mini_card(s, Inches(4.73), Inches(3.78), Inches(3.83), Inches(1.90), "The brief is the product",
              "The Blog Agent has no web search and no brochures — it can only write from the brief. "
              "That makes the Research Agent's brief the single thing worth checking.", TEAL, 12, 10)
    mini_card(s, Inches(8.74), Inches(3.78), Inches(3.83), Inches(1.90), "Probe the approval",
              "The Manager is instructed to stop for approval — a rule, not a gate. The test script "
              "tries to talk it out of the review. Contrast Lab 8, where the platform suspends the run.",
              RED, 12, 10)

    lab_banner(s, "4B", "Lab 4b — Multi-Agent Content Team",
               "Research → draft → review → human approval  ·  Copilot Studio connected agents  ·  60–75 min · optional")


def draw_lab5b(s):
    text(s, Inches(0.72), Inches(1.66), Inches(11.85), Inches(0.34),
         "The mirror image of Lab 5: the agent decides when to call the workflow, and the flow's inputs and "
         "outputs are the contract it programs against.", 13, GREY, anchor=MSO_ANCHOR.TOP)
    if FLOW_5B.exists():
        with Image.open(FLOW_5B) as im:
            iw, ih = im.size
        max_w, max_h = Inches(12.0), Inches(3.62)
        scale = min(max_w / iw, max_h / ih)
        w, h = int(iw * scale), int(ih * scale)
        top = Inches(2.12) + int((Inches(3.62) - h) / 2)
        s.shapes.add_picture(str(FLOW_5B), int((SW - w) / 2), top, w, h)
    lab_banner(s, "5B", "Lab 5b — Calling Workflow from Agent",
               "Blog Writer Agent runs the flow as its tool  ·  Copilot Studio + M365 Copilot  ·  30 min")


def draw_lab_journey(s):
    rows = [
        ("0", "Environment Setup", "Developer environment, Dataverse, credits", BLUE),
        ("1", "Trigger and Actions", "Form → email; dynamic content", BLUE),
        ("2", "Log to Excel", "Commit the audit row before confirming", BLUE),
        ("3", "Leave Application Approval", "The flow pauses for a manager", TEAL),
        ("4", "Agents", "Instructions · skills · knowledge · tools", VIOLET),
        ("4b", "Multi-Agent Content Team", "Research → draft → review → human approval", VIOLET),
        ("5", "Calling Agent from Workflow", "Topic in, M365 Copilot blog, posted to Teams", VIOLET),
        ("5b", "Calling Workflow from Agent", "The agent runs the flow as its tool", VIOLET),
        ("6", "HTTP + Approval Agent", "Structured output; the AI never writes the record", GREEN),
        ("7", "HTTP + Chatbot", "The agent alone, in public", GREEN),
        ("8", "HTTP + Human Review", "The run that will not proceed", RED),
        ("9", "RAG with Knowledge Base", "Three nodes, no ingestion", AMBER),
        ("10", "RAG with Pinecone", "Chunking, embeddings, top_k", AMBER),
    ]
    for i, (num, name, note, color) in enumerate(rows):
        col, row = (0, i) if i < 7 else (1, i - 7)
        x = Inches(0.72 + col * 6.0)
        y = Inches(1.62 + row * 0.63)
        box(s, x, y, Inches(5.85), Inches(0.57), LIGHT, LINE)
        chip = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.10), y + Inches(0.07), Inches(0.44), Inches(0.44))
        chip.fill.solid()
        chip.fill.fore_color.rgb = color
        chip.line.fill.background()
        chip.shadow.inherit = False
        text(s, x + Inches(0.10), y + Inches(0.07), Inches(0.44), Inches(0.44), num, 11.5, WHITE, True, PP_ALIGN.CENTER)
        text(s, x + Inches(0.66), y + Inches(0.005), Inches(5.08), Inches(0.30), name, 11.5, INK, True)
        text(s, x + Inches(0.66), y + Inches(0.285), Inches(5.08), Inches(0.27), note, 9, GREY)
    takeaway(s, 6.14, "Thirteen labs, each reusing the one before it. Nothing is rebuilt from scratch.")


# --------------------------------------------------------------- helpers

def slide_heading(s):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            if sh.top is not None and Inches(0.5) < sh.top < Inches(1.2):
                return sh.text_frame.text.strip().split("\n")[0]
    return ""


def replace_text(slide_obj, old, new, required=True):
    hits = 0
    for sh in slide_obj.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    hits += 1
            if hits:
                continue
            joined = "".join(r.text for r in para.runs)
            if old in joined and para.runs:
                para.runs[0].text = joined.replace(old, new)
                for r in para.runs[1:]:
                    r.text = ""
                hits += 1
    if required and hits == 0:
        raise SystemExit(f"no match for {old[:60]!r}")
    return hits


# --------------------------------------------------------------- apply

# 1. collision-safe partnames
_used_nums = set()
for _s in prs.slides:
    _m = re.search(r"slide(\d+)\.xml$", str(_s.part.partname))
    if _m:
        _used_nums.add(int(_m.group(1)))


def _next_free_slide_partname(self):
    n = max(_used_nums) + 1
    _used_nums.add(n)
    return PackURI("/ppt/slides/slide%d.xml" % n)


PresentationPart._next_slide_partname = property(_next_free_slide_partname)

before = len(prs.slides)
id_list = prs.slides._sldIdLst

# 2. rebuild the journey slide (slide 8) for 13 labs
s8 = prs.slides[7]
assert "Eleven Connected Labs" in "".join(
    sh.text_frame.text for sh in s8.shapes if sh.has_text_frame)
for sh in list(s8.shapes):
    sh._element.getparent().remove(sh._element)
title(s8, "Thirteen Connected Labs", "the journey")
draw_lab_journey(s8)
footer(s8, 8)
print("slide 8 rebuilt for 13 labs")

# 3. schedule rows (slide 10)
s10 = prs.slides[9]
replace_text(s10, "Lab 3 · Module 3 · Lab 4 — approval, agent anatomy, four agents",
             "Lab 3 · Module 3 · Lab 4 · 4b — approval, agents, multi-agent team")
replace_text(s10, "Recap · Lab 5 — calling an agent from a workflow",
             "Recap · Lab 5 · 5b — agent and workflow, each calling the other")
print("slide 10 schedule updated")

# 4. Module 3 divider applied-in line (slide 31)
s31 = prs.slides[30]
replace_text(s31, "Applied in Lab 4 and Lab 5", "Applied in Labs 4, 4b, 5 and 5b", required=False) or \
    replace_text(s31, "Lab 4 and Lab 5", "Labs 4, 4b, 5 and 5b", required=False)
print("slide 31 divider updated")

# 5. cover version
replace_text(prs.slides[0], "Version 7.2", "Version 7.3")
print("cover bumped to 7.3")

# 6. insert the four new slides
NEW = [
    ("Agent Settings", "module 3 · concept", draw_agent_settings,
     "Multiple Agents — A Split Is a Governance Decision"),
    ("Publishing an Agent — Channels", "module 3 · concept", draw_publishing_channels,
     "Publishing to Teams"),
    ("Lab 4b — Multi-Agent Content Team", "module 3 · lab workflow", draw_lab4b,
     "Lab 4 — Four Agents, One Spine"),
    ("Lab 5b — Calling Workflow from Agent", "module 3 · lab workflow", draw_lab5b,
     "Lab 5 — Calling Agent from Workflow"),
]
for heading, kicker, draw, anchor_title in NEW:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, heading, kicker)
    draw(s)
    footer(s, 0)
    el = list(id_list)[-1]
    pos = None
    for i, existing in enumerate(prs.slides):
        if existing is s:
            continue
        if slide_heading(existing) == anchor_title:
            pos = i + 1
            break
    assert pos is not None, f"anchor not found: {anchor_title}"
    id_list.remove(el)
    id_list.insert(pos, el)
    print(f"inserted {heading!r} at position {pos + 1}")

assert len(prs.slides) == before + 4

# 7. renumber footers
for no, sl in enumerate(prs.slides, 1):
    for sh in sl.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t.isdigit() and sh.top is not None and sh.top > Inches(6.6) and (sh.left or 0) > Inches(0.5):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.strip().isdigit():
                        r.text = str(no)

prs.save(DECK)

# 8. clean duplicate zip entries
tmp = DECK.with_suffix(".pptx.repaired")
with zipfile.ZipFile(DECK) as zin:
    last = {}
    for idx, info in enumerate(zin.infolist()):
        last[info.filename] = idx
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for idx, info in enumerate(zin.infolist()):
            if last[info.filename] == idx:
                zout.writestr(info, zin.read(info))
shutil.move(tmp, DECK)

# 9. archive the superseded deck + pdf
archive = CW / "archive"
archive.mkdir(exist_ok=True)
for stale in (SRC, CW / f"{TITLE}-v7.2.pdf"):
    if stale.exists():
        shutil.move(str(stale), archive / stale.name)
        print("archived", stale.name)

print(f"Saved {DECK.name} with {len(prs.slides)} slides")
for i, s in enumerate(prs.slides, 1):
    print(f"{i:3d}  {slide_heading(s)[:70]}")
