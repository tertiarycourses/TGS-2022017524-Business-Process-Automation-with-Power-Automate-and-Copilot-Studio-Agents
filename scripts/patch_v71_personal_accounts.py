#!/usr/bin/env python3
"""Insert a "Microsoft 365 Personal Accounts" slide after the training-account
(Copilot sign-in) slide.

Adds the two shared outlook.com personal accounts (training1-tertiary /
training2-tertiary, password Tertiary@888) as a house-style slide directly
after "Your Training Account", then renumbers every footer page number.
Edits the v7.1 deck in place; the pre-patch copy is archived first by the
calling shell step.
"""

from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.packuri import PackURI
from pptx.parts.presentation import PresentationPart
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
CW = ROOT / "courseware"
DECK = CW / "Business Process Automation with Power Automate and Copilot Studio Agents-v7.1.pptx"

BLUE = RGBColor(0x1F, 0x6F, 0xEB)
VIOLET = RGBColor(0x6D, 0x3F, 0xD2)
TEAL = RGBColor(0x10, 0x8A, 0x73)
INK = RGBColor(0x16, 0x1B, 0x26)
GREY = RGBColor(0x5B, 0x63, 0x72)
LINE = RGBColor(0xD7, 0xE0, 0xEA)
LIGHT = RGBColor(0xF5, 0xF8, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FOOTER_TEXT = (
    "Business Process Automation with Power Automate & Copilot Studio Agents  ·  "
    "TGS-2022017524  ·  © 2026 Tertiary Infotech Academy Pte Ltd"
)

prs = Presentation(DECK)


def box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True, line_w=1.2):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_w)
    shape.shadow.inherit = False
    return shape


def text(slide, x, y, w, h, value, size=18, color=INK, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=0.08, space=0):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    for i, line_text in enumerate(str(value).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.space_after = Pt(space)
        r = p.add_run()
        r.text = line_text
        r.font.name = "Arial"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return tb


def title(slide, heading, kicker):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), Inches(1.48))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    bar.shadow.inherit = False
    text(slide, Inches(0.72), Inches(0.34), Inches(11.8), Inches(0.28), kicker.upper(), 12, BLUE, True)
    text(slide, Inches(0.72), Inches(0.68), Inches(11.8), Inches(0.66), heading, 29, INK, True)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.5), Inches(11.85), Inches(0.015))
    rule.fill.solid()
    rule.fill.fore_color.rgb = LINE
    rule.line.fill.background()
    rule.shadow.inherit = False


def footer(slide, slide_no):
    text(slide, Inches(0.72), Inches(6.95), Inches(11.85), Inches(0.24), FOOTER_TEXT, 8, GREY)
    text(slide, Inches(0.72), Inches(6.95), Inches(11.85), Inches(0.24), str(slide_no), 8, GREY, False, PP_ALIGN.RIGHT)


def takeaway(slide, y, message, color=BLUE):
    box(slide, Inches(0.72), Inches(y), Inches(11.85), Inches(0.62), LIGHT, color, line_w=1.6)
    text(slide, Inches(1.0), Inches(y), Inches(11.3), Inches(0.62), message, 14, INK, True)


def draw_personal_accounts(s):
    text(s, Inches(0.72), Inches(1.62), Inches(11.85), Inches(0.34),
         "Two shared Microsoft 365 personal accounts — use these when a lab step asks for a personal "
         "(outlook.com) sign-in.", 13, GREY, anchor=MSO_ANCHOR.TOP)
    accounts = [
        ("training1-tertiary@outlook.com", "Personal account 1"),
        ("training2-tertiary@outlook.com", "Personal account 2"),
    ]
    for i, (email, who) in enumerate(accounts):
        x = Inches(0.72 + i * 6.05)
        y = Inches(2.3)
        box(s, x, y, Inches(5.8), Inches(0.9), WHITE, LINE)
        text(s, x + Inches(0.24), y, Inches(3.9), Inches(0.9), email, 15, INK, True)
        text(s, x + Inches(3.9), y, Inches(1.75), Inches(0.9), who, 11, GREY, False, PP_ALIGN.RIGHT)
    takeaway(s, 3.7, "Password for both accounts:  Tertiary@888", TEAL)
    text(s, Inches(0.72), Inches(4.6), Inches(11.85), Inches(0.6),
         "These are shared personal accounts, separate from your assigned "
         "training account on the tertiaryinfotech tenant — sign in with them only where a lab calls for a "
         "personal Microsoft account, and do not change the password.", 11, GREY, anchor=MSO_ANCHOR.TOP)


# --- collision-safe partname allocation (partnames in this file have gaps) ---
_used_nums = set()
for _s in prs.slides:
    _m = re.search(r"slide(\d+)\.xml$", str(_s.part.partname))
    if _m:
        _used_nums.add(int(_m.group(1)))


def _next_free_slide_partname(self):
    n = max(_used_nums) + 1
    _used_nums.add(n)
    return PackURI("/ppt/slides/slide%d.xml" % n)


PresentationPart._next_slide_partname = property(_next_free_slide_partname)

# --- build and position the slide -------------------------------------------
before = len(prs.slides)
id_list = prs.slides._sldIdLst

s = prs.slides.add_slide(prs.slide_layouts[6])
title(s, "Microsoft 365 Personal Accounts", "course environment")
draw_personal_accounts(s)
footer(s, 0)  # renumbered below

anchor_pos = None
for i, existing in enumerate(prs.slides):
    if existing is s:
        continue
    for sh in existing.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == "Your Training Account":
            anchor_pos = i + 1
            break
    if anchor_pos is not None:
        break
assert anchor_pos is not None, "anchor slide 'Your Training Account' not found"

el = list(id_list)[-1]
id_list.remove(el)
id_list.insert(anchor_pos, el)
assert len(prs.slides) == before + 1

# --- renumber every footer page number ---------------------------------------
for no, sl in enumerate(prs.slides, 1):
    for sh in sl.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t.isdigit() and sh.top is not None and sh.top > Inches(6.6) and (sh.left or 0) > Inches(0.5):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.strip().isdigit():
                        r.text = str(no)

prs.save(DECK)

# keep only the last occurrence of each zip entry name (clean package)
tmp = DECK.with_suffix(".pptx.repaired")
with zipfile.ZipFile(DECK) as zin:
    last = {}
    for idx, info in enumerate(zin.infolist()):
        last[info.filename] = idx
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for idx, info in enumerate(zin.infolist()):
            if last[info.filename] == idx:
                zout.writestr(info, zin.read(info))
shutil.move(tmp, DECK)

print(f"Saved {DECK.name} with {len(prs.slides)} slides (was {before})")
print(f"New slide inserted at position {anchor_pos + 1}")
