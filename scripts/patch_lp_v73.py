#!/usr/bin/env python3
"""LP v7.3 — align the Lesson Plan to the 13-lab structure and the v7.3 deck.

Day 2 is retimed to give Lab 5b its 30 minutes (recap 10→5, Module 4 30→20,
Lab 6 60→50, Module 5 15→10; every day still totals 8 scheduled hours after a
1-hour lunch). Lab 4b enters Day 1 as the optional extension of Lab 4. All
Slides mappings move to the 80-slide v7.3 deck, and the Lab Alignment Reference
gains Lab 4b and Lab 5b entries. The deck's own condensed schedule slide (10)
is retimed to match.
"""

from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
LP = ROOT / "courseware" / "LP-Business Process Automation with Power Automate and Copilot Studio Agents.docx"

doc = Document(LP)

# 1. version header + deck pointer
for p in doc.paragraphs:
    if p.text.strip() == "Version 7.2":
        p.runs[0].text = "Version 7.3"
        break
for p in doc.paragraphs:
    if "-v7.2.pptx, 76 slides" in p.text:
        for r in p.runs:
            if "v7.2" in r.text or "76 slides" in r.text:
                r.text = r.text.replace("v7.2", "v7.3").replace("76 slides", "80 slides")
        assert "v7.3" in p.text and "80 slides" in p.text
        break

# 2. version control record row
vt = doc.tables[0]
row = vt.add_row()
for ci, val in enumerate([
    "7.3", "6 Aug 2026",
    "Aligned to the 13-lab structure and the 80-slide v7.3 deck — Lab 4b (Multi-Agent "
    "Content Team) added as the optional Day 1 extension of Lab 4, Lab 5b (Calling "
    "Workflow from Agent) scheduled after Lab 5 on Day 2, and all slide mappings updated.",
    "Course Development Team",
]):
    row.cells[ci].paragraphs[0].text = val
    for r in row.cells[ci].paragraphs[0].runs:
        r.font.size = Pt(9)

# 3. overview paragraphs + Labs summary row
for p in doc.paragraphs:
    if p.text.endswith("agent anatomy — Labs 0–4."):
        p.runs[-1].text = p.runs[-1].text.replace("Labs 0–4.", "Labs 0–4, with the optional Lab 4b.")
        break
for p in doc.paragraphs:
    if "RAG twice over — Labs 5–10" in p.text:
        for r in p.runs:
            if "Labs 5–10" in r.text:
                r.text = r.text.replace("Labs 5–10", "Labs 5, 5b and 6–10")
        break
labs_cell = doc.tables[1].rows[6].cells[1]
labs_cell.paragraphs[0].text = (
    "13 step-by-step labs across 2 days (Lab 0 setup plus Labs 1–10, plus the optional "
    "Lab 4b and Lab 5b: Labs 0–4b on Day 1, Labs 5–10 on Day 2), driven by 5 concept modules"
)

# 4. Day 1 schedule (table 2): Module 3 span, Lab 4 row + optional 4b, recap span
day1 = doc.tables[2]


def set_row(table, ri, values):
    for ci, val in enumerate(values):
        if val is None:
            continue
        cell = table.rows[ri].cells[ci]
        keep_size = None
        for r in cell.paragraphs[0].runs:
            if keep_size is None:
                keep_size = r.font.size
        cell.paragraphs[0].text = val
        for r in cell.paragraphs[0].runs:
            r.font.size = keep_size or Pt(9.5)


for ri, r in enumerate(day1.rows):
    cells = [c.text.strip() for c in r.cells]
    if cells[2].startswith("Module 3:"):
        set_row(day1, ri, [None, None,
                           "Module 3: workflow versus agent; the new Copilot Studio (model and harness, "
                           "the new agent designer); agent anatomy — instructions, skills, knowledge, tools; "
                           "what is actually enforced; connected agents; agent settings; publishing and channels",
                           None, "31–46"])
    elif cells[2].startswith("Lab 4 —"):
        set_row(day1, ri, [None, None,
                           "Lab 4 — Agents — assemble the Procurement, HR, Sales and IT Support agents "
                           "(optional extension: Lab 4b — Multi-Agent Content Team, 60–75 min, self-paced)",
                           None, "47–48"])
    elif cells[2].startswith("Day 1 recap"):
        set_row(day1, ri, [None, None, None, None, "31–46"])

# 5. Day 2 schedule (table 3): full retime with Lab 5b
day2 = doc.tables[3]
NEW_DAY2 = [
    ("9:00 – 9:05", "5 min", "Day 1 recap and Q&A", "Facilitated discussion", "31–46"),
    ("9:05 – 9:35", "30 min",
     "Lab 5 — Calling Agent from Workflow — Start-node topic input, M365 Copilot drafts "
     "the blog, posted to the Training team's General channel", "Hands-on lab", "49"),
    ("9:35 – 10:05", "30 min",
     "Lab 5b — Calling Workflow from Agent — the same boundary crossed the other way: the "
     "Blog Writer Agent runs the published flow as its tool", "Hands-on lab", "50"),
    ("10:05 – 10:25", "20 min",
     "Module 4: agent flows and the Agent node, the new workflow designer, HTTP request "
     "and response, JSON schema, structured output, the boundary of agency", "Lecture & demo", "51–58"),
    ("10:25 – 10:40", "15 min", "Tea break", "—", "—"),
    ("10:40 – 11:30", "50 min",
     "Lab 6 — HTTP and Application Approval Agent — web form to agent flow, six ordered "
     "rules, four decisions", "Hands-on lab", "59"),
    ("11:30 – 12:10", "40 min",
     "Lab 7 — HTTP and Chatbot — the agent alone in public, and the probes that break it",
     "Hands-on lab", "60"),
    ("12:10 – 1:10", "60 min", "Lunch", "—", "—"),
    ("1:10 – 2:10", "60 min",
     "Lab 8 — HTTP and Human Review — the Human review gate and the run that will not proceed",
     "Hands-on lab", "61"),
    ("2:10 – 2:20", "10 min",
     "Module 5: why retrieval rather than a bigger prompt; the RAG pipeline; chunking, "
     "embeddings and top_k", "Lecture & demo", "62–66"),
    ("2:20 – 3:00", "40 min",
     "Lab 9 — RAG with Knowledge Base — RAG as a product setting: three nodes, no ingestion",
     "Hands-on lab", "67"),
    ("3:00 – 3:15", "15 min", "Tea break", "—", "—"),
    ("3:15 – 3:55", "40 min",
     "Lab 10 — RAG with Pinecone — rebuild the same chatbot on Pinecone and compare the levers",
     "Hands-on lab", "68"),
    ("3:55 – 4:00", "5 min", "Course synthesis and briefing before the assessment",
     "Facilitated discussion", "69–72"),
    ("4:00 – 5:00", "60 min", "Written Assessment (WA / SAQ) — 1 hour, open book",
     "Individual assessment", "73–80"),
    ("5:00 – 6:00", "60 min", "Practical Performance (PP) Assessment — 1 hour, open book",
     "Individual assessment", "73–80"),
]
# grow the table by one data row (clone the last row's XML for identical formatting)
while len(day2.rows) < len(NEW_DAY2) + 1:
    clone = copy.deepcopy(day2.rows[-1]._tr)
    day2.rows[-1]._tr.addnext(clone)
assert len(day2.rows) == len(NEW_DAY2) + 1
for ri, values in enumerate(NEW_DAY2, 1):
    set_row(day2, ri, list(values))

# 6. Lab Alignment Reference — renumber and insert 4b / 5b
RENUMBER = {
    "Lab 4 — Agents · Slides 45": "Lab 4 — Agents · Slides 47",
    "Lab 5 — Calling Agent from Workflow · Slides 46": "Lab 5 — Calling Agent from Workflow · Slides 49",
    "Lab 6 — HTTP and Application Approval Agent · Slides 55": "Lab 6 — HTTP and Application Approval Agent · Slides 59",
    "Lab 7 — HTTP and Chatbot · Slides 56": "Lab 7 — HTTP and Chatbot · Slides 60",
    "Lab 8 — HTTP and Human Review · Slides 57": "Lab 8 — HTTP and Human Review · Slides 61",
    "Lab 9 — RAG with Knowledge Base · Slides 63": "Lab 9 — RAG with Knowledge Base · Slides 67",
    "Lab 10 — RAG with Pinecone · Slides 64": "Lab 10 — RAG with Pinecone · Slides 68",
}
for p in doc.paragraphs:
    if p.style.name == "Heading 3" and p.text.strip() in RENUMBER:
        new = RENUMBER[p.text.strip()]
        for r in p.runs:
            r.text = ""
        (p.runs[0] if p.runs else p.add_run()).text = new


def insert_h3_with_body(after_body_para, h3_text, body_text, template_h3, template_body):
    h3 = copy.deepcopy(template_h3._p)
    body = copy.deepcopy(template_body._p)
    after_body_para._p.addnext(body)
    after_body_para._p.addnext(h3)
    hp, bp = Paragraph(h3, after_body_para._parent), Paragraph(body, after_body_para._parent)
    for para, txt in ((hp, h3_text), (bp, body_text)):
        for r in para.runs:
            r.text = ""
        (para.runs[0] if para.runs else para.add_run()).text = txt


paras = doc.paragraphs
for i, p in enumerate(paras):
    if p.style.name == "Heading 3" and p.text.startswith("Lab 4 — Agents"):
        insert_h3_with_body(paras[i + 1],
                            "Lab 4b — Multi-Agent Content Team · Slides 48",
                            "Duration: 60–75 minutes (optional — extends Lab 4). Follow the matching "
                            "detailed, step-by-step lab and its shared workflow visual in the Learner Guide.",
                            p, paras[i + 1])
        break
paras = doc.paragraphs
for i, p in enumerate(paras):
    if p.style.name == "Heading 3" and p.text.startswith("Lab 5 — Calling Agent"):
        insert_h3_with_body(paras[i + 1],
                            "Lab 5b — Calling Workflow from Agent · Slides 50",
                            "Duration: 30 minutes. Follow the matching detailed, step-by-step lab and its "
                            "shared workflow visual in the Learner Guide.",
                            p, paras[i + 1])
        break

doc.save(LP)
print("LP saved")

# ---- deck slide 10: retime the condensed Day 2 rows to match the LP ----
from pptx import Presentation

DECK = ROOT / "courseware" / "Business Process Automation with Power Automate and Copilot Studio Agents-v7.3.pptx"
prs = Presentation(DECK)
s10 = prs.slides[9]
TIME_MAP = {
    "9:00 – 9:40": "9:00 – 10:05",
    "9:40 – 11:10": "10:05 – 11:30",
    "11:25 – 12:05": "11:30 – 12:10",
    "1:05 – 2:05": "1:10 – 2:10",
    "2:05 – 3:55": "2:10 – 3:55",
}
hits = 0
for sh in s10.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text.strip()
    if t in TIME_MAP:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.text.strip() == t:
                    r.text = TIME_MAP[t]
                    hits += 1
assert hits == len(TIME_MAP), f"only {hits} time cells updated"
prs.save(DECK)
print("deck slide 10 retimed")
